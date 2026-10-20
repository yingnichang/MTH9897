"""Build the final PDF slides from the same data-driven content as build_deck.py.

Project authors: Yingni Chang and Yikai Shen.

Run: python presentation/build_pdf.py
Requires python-pptx, reportlab and Pillow. Text and tables remain searchable.
"""
from io import BytesIO
from pathlib import Path
from types import SimpleNamespace
from xml.sax.saxutils import escape

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from reportlab.lib.colors import HexColor, white
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfgen import canvas
from reportlab.platypus import Paragraph
from reportlab.lib.utils import ImageReader
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont

import sys
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
from research import results
from research.feedback_summary import extras
from numeric_views import make_views
from chart_theme import chart as themed_chart

EMU = 12700
FONT, BOLD = 'Helvetica', 'Helvetica-Bold'
font_dir = Path('C:/Windows/Fonts')
if (font_dir / 'arial.ttf').exists():
    pdfmetrics.registerFont(TTFont('SlideArial', str(font_dir / 'arial.ttf')))
    pdfmetrics.registerFont(TTFont('SlideArialBold', str(font_dir / 'arialbd.ttf')))
    FONT, BOLD = 'SlideArial', 'SlideArialBold'
else:
    import matplotlib
    font_dir = Path(matplotlib.get_data_path()) / 'fonts/ttf'
    pdfmetrics.registerFont(TTFont('SlideSans', str(font_dir / 'DejaVuSans.ttf')))
    pdfmetrics.registerFont(TTFont('SlideSansBold', str(font_dir / 'DejaVuSans-Bold.ttf')))
    FONT, BOLD = 'SlideSans', 'SlideSansBold'


NAVY, TEAL, CREAM = HexColor('#102C3B'), HexColor('#087F83'), HexColor('#F5F3EE')
INK, MUTED = HexColor('#183342'), HexColor('#536B73')
HEADER, RULE, STRIPE = NAVY, HexColor('#D5DEDA'), HexColor('#E7EFEB')


def paragraph(pdf, text, x, top, width, size, bold=False, color=INK, align=0, measure=False):
    style = ParagraphStyle('slide', fontName=BOLD if bold else FONT,
                           fontSize=size, leading=size * 1.28, textColor=color,
                           alignment=align, splitLongWords=False)
    p = Paragraph(escape(text).replace('\n', '<br/>'), style)
    _, height = p.wrap(width, 10000)
    if not measure:
        p.drawOn(pdf, x, top - height)
    return height


def body_height(paras, width, size, gap):
    return sum(paragraph(None, p.text, 0, 0, width-(18 if p.level else 0),
                         size-(2 if p.level else 0), measure=True)
               for p in paras) + gap * max(0, len(paras)-1)


def fit_body(paras, width, height, preferred, minimum=16):
    for half in range(round(preferred*2), round(minimum*2)-1, -1):
        size = half/2
        gap = max(8, size*.48)
        used = body_height(paras, width, size, gap)
        if used <= height:
            return size, gap, used
    raise ValueError('Slide body does not fit without changing content')


def draw_body(pdf, paras, x, top, width, size, gap):
    for i, p in enumerate(paras):
        indent = 18 if p.level else 0
        used = paragraph(pdf, p.text, x+indent, top, width-indent,
                         size-(2 if p.level else 0), color=MUTED if p.level else INK)
        top -= used + (gap if i+1 < len(paras) else 0)
    return top


def table_layout(table, width, size, fractions):
    widths = [width*f for f in fractions]
    cells = table if isinstance(table, list) else [[c.text for c in row.cells] for row in table.rows]
    heights = []
    for row_i, row in enumerate(cells):
        needed = max(paragraph(None, value, 0, 0, cw-18, size,
                               bold=row_i == 0 or col_i == 0, measure=True)
                     for col_i, (value, cw) in enumerate(zip(row, widths)))
        heights.append(max(size*1.95, needed+10))
    return widths, heights, cells


def fit_table(table, width, height, preferred, fractions):
    for half in range(round(preferred*2), 25, -1):
        size = half/2
        layout = table_layout(table, width, size, fractions)
        words_fit = all(pdfmetrics.stringWidth(word, BOLD if ri == 0 or ci == 0 else FONT, size)
                        <= layout[0][ci]-18
                        for ri, row in enumerate(layout[2]) for ci, value in enumerate(row)
                        for word in value.split())
        if words_fit and sum(layout[1]) <= height:
            return size, layout
    raise ValueError('Slide table does not fit without changing content')


def draw_table(pdf, x, top, size, layout, prose=False):
    widths, heights, cells = layout
    for row_i, (row, rh) in enumerate(zip(cells, heights)):
        pdf.setFillColor(HEADER if row_i == 0 else STRIPE if row_i % 2 else white)
        pdf.rect(x, top-rh, sum(widths), rh, stroke=0, fill=1)
        left = x
        for col_i, (value, cw) in enumerate(zip(row, widths)):
            align = 0 if col_i == 0 or prose else 1
            needed = paragraph(None, value, 0, 0, cw-18, size,
                               bold=row_i == 0 or col_i == 0, measure=True)
            paragraph(pdf, value, left+9, top-(rh-needed)/2, cw-18, size,
                      bold=row_i == 0 or col_i == 0,
                      color=white if row_i == 0 else INK, align=align)
            left += cw
        pdf.setStrokeColor(RULE)
        pdf.setLineWidth(.6)
        pdf.line(x, top-rh, x+sum(widths), top-rh)
        top -= rh
    return top


def draw_picture(pdf, shape, x, top, width, height):
    ratio = shape.width/shape.height
    w = min(width, height*ratio)
    h = w/ratio
    pdf.drawImage(ImageReader(BytesIO(shape.image.blob)), x+(width-w)/2,
                  top-(height-h)/2-h, w, h, mask='auto')


def text_paras(lines):
    return [SimpleNamespace(text=line, level=0) for line in lines]


def draw_math(pdf, expression, x, top, width, height, size=19):
    """Render mathematical glyphs and rules as searchable vector PDF content."""
    from matplotlib.mathtext import MathTextParser
    from matplotlib.font_manager import FontProperties
    parsed=MathTextParser('path').parse(expression,dpi=72,prop=FontProperties(size=size))
    scale=min(1,width/parsed.width,height/parsed.height)
    pdf.saveState()
    pdf.translate(x+(width-parsed.width*scale)/2,top-(parsed.height-parsed.depth)*scale)
    pdf.scale(scale,scale);pdf.setFillColor(INK)
    for glyph in parsed.glyphs:
        # Matplotlib 3.11 adds the glyph index after the Unicode codepoint.
        if len(glyph)==6:
            font,font_size,code,_,gx,gy=glyph
        else:
            font,font_size,code,gx,gy=glyph
        name='SlideMath_'+font.postscript_name
        if name not in pdfmetrics.getRegisteredFontNames():
            pdfmetrics.registerFont(TTFont(name,font.fname))
        pdf.setFont(name,font_size);pdf.drawString(gx,gy,chr(code))
    for rx,ry,rw,rh in parsed.rects:
        pdf.rect(rx,ry,rw,rh,stroke=0,fill=1)
    pdf.restoreState()


def research_questions(pdf,x,top,width,height):
    """The three research questions, in the same order as both speaking scripts."""
    gap=34
    card_width=(width-2*gap)/3
    card_top=top-8
    card_height=258
    cards=[
        ('01  REPLICATE','Can we reproduce the paper?',
         '1929–2016\nReturn, volatility and turnover'),
        ('02  EXPLAIN','How much return comes from known factors?',
         'Carhart: primary model\nFF5 + momentum: robustness'),
        ('03  TEST RECENT PERFORMANCE','Does it still perform after trading costs?',
         'Recent periods through 2024\n10 / 30 bps and break-even costs\n2017–2024: 0–30 bps grid'),
    ]
    for i,(label,question,detail) in enumerate(cards):
        left=x+i*(card_width+gap)
        pdf.setFillColor(STRIPE if i==1 else white);pdf.setStrokeColor(RULE)
        pdf.setLineWidth(.8)
        pdf.roundRect(left,card_top-card_height,card_width,card_height,12,stroke=1,fill=1)
        pdf.setFillColor(TEAL);pdf.roundRect(left+18,card_top-37,30,4,2,stroke=0,fill=1)
        paragraph(pdf,label,left+18,card_top-49,card_width-36,10,bold=True,color=TEAL)
        paragraph(pdf,question,left+18,card_top-83,card_width-36,22,bold=True)
        paragraph(pdf,detail,left+18,card_top-183,card_width-36,15.5,color=MUTED)
        if i<2:
            cy=card_top-card_height/2
            start=left+card_width+7;end=start+gap-14
            pdf.setStrokeColor(TEAL);pdf.setLineWidth(1.8)
            pdf.line(start,cy,end,cy)
            pdf.line(end-5,cy+4,end,cy);pdf.line(end-5,cy-4,end,cy)
    bottom_top=card_top-card_height-22
    paragraph(pdf,'Time checks: 1929–1962 | 1963–1989 | 1990–2016 | 2017–2024; rolling 120-month windows',
              x,bottom_top,width,13.5,color=MUTED)
    paragraph(pdf,'Uncertainty: Newey–West (12 lags) or block-bootstrap confidence intervals',
              x,bottom_top-27,width,13.5,color=MUTED)


def selection_flow(pdf, paras, x, top, width, height):
    """Vector selection diagram, with the original implementation conventions alongside."""
    notes = paras[1:]
    note_size, note_gap, note_height = fit_body(notes, width, 85, 12, minimum=11)
    main_height = height-note_height-15
    flow_width = width*.53
    right_x = x+flow_width+22
    right_width = width-flow_width-22
    scale = main_height/312

    def arrow(offset, length=13):
        cx = x+flow_width/2
        y = top-offset*scale
        end = y-length*scale
        pdf.setStrokeColor(TEAL);pdf.setLineWidth(1.5)
        pdf.line(cx,y,cx,end)
        pdf.line(cx-4,end+5,cx,end);pdf.line(cx+4,end+5,cx,end)

    def node(left, offset, node_width, node_height, title, detail, fill, title_size=19, detail_size=13,
             equation=None, equation_height=24):
        node_top = top-offset*scale
        node_height *= scale
        pdf.setFillColor(fill);pdf.setStrokeColor(RULE);pdf.setLineWidth(.7)
        pdf.roundRect(left,node_top-node_height,node_width,node_height,10,stroke=1,fill=1)
        if equation:
            paragraph(pdf,title,left+10,node_top-6,node_width-20,title_size,bold=True,align=1)
            draw_math(pdf,equation,left+10,node_top-34,node_width-20,equation_height)
            return
        title_h = paragraph(None,title,0,0,node_width-20,title_size,bold=True,measure=True)
        detail_h = paragraph(None,detail,0,0,node_width-20,detail_size,measure=True)
        text_top = node_top-(node_height-title_h-detail_h-5)/2
        paragraph(pdf,title,left+10,text_top,node_width-20,title_size,bold=True,align=1)
        paragraph(pdf,detail,left+10,text_top-title_h-5,node_width-20,detail_size,color=MUTED,align=1)

    node(x,0,flow_width,48,'1,000 largest eligible US stocks',
         '36-month history screen, then market capitalization',white,18,13)
    arrow(52,9)
    node(x,66,flow_width,68,'500 lower-volatility stocks','',STRIPE,19,13,
         equation=r'$\sigma_{i,t}=\sqrt{12}\times\mathrm{Std}(R_{i,t-35},\ldots,R_{i,t})$')
    arrow(138,9)
    pair_width=(flow_width-12)/2
    node(x,152,pair_width,78,'Momentum','',white,18,12,
         equation=r'$MOM_{i,t}=\prod_{j=1}^{11}(1+r_{i,t-j}^{\mathrm{price}})-1$',equation_height=34)
    node(x+pair_width+12,152,pair_width,78,'Net payout yield',
         'Dividends + net share repurchases',white,18,12)
    paragraph(pdf,'Average the two ranks; select the best 100',x,top-234*scale,
              flow_width,12,bold=True,align=1,color=TEAL)
    arrow(252,7)
    node(x,264,flow_width,48,'100 stocks',
         'Equal-weighted portfolio; rebalance quarterly',HexColor('#DCECE9'),18,12)

    conventions = [
        ['Convention', 'Implementation'],
        ['Share classes', 'Largest common class per company each month; share codes 10/11.'],
        ['Size measure', 'Security market capitalization after the history screen.'],
        ['Dividends', 'Implied from RET - RETX; legacy RETX includes non-ordinary distributions.'],
        ['Missing signals', 'Partition by volatility first, then rank complete cases.'],
        ['Early universe', 'Use all eligible names if fewer than 1,000: 393 at 1928-12; full 500-name pool from 1964-06.'],
        ['Missing returns', 'Stale carry (0) or an in-gap price ratio.'],
    ]
    size,layout=fit_table(conventions,right_width,main_height,14,[.32,.68])
    draw_table(pdf,right_x,top,size,layout,prose=True)
    draw_body(pdf,notes,x,top-main_height-15,width,note_size,note_gap)


def draw_numeric_view(pdf, view, pictures, tables, x, top, width, height):
    """Place a verified numerical table with its context and source figure."""
    kind = view['kind']
    intro, notes = text_paras(view['intro']), text_paras(view['notes'])
    bottom = top-height
    if kind == 'payout_formula':
        pdf.setFillColor(STRIPE);pdf.setStrokeColor(RULE);pdf.setLineWidth(.7)
        pdf.roundRect(x,top-108,width,108,11,stroke=1,fill=1)
        paragraph(pdf,'Net payout yield = dividend yield + net share repurchase proxy',
                  x+18,top-10,width-36,17,bold=True,color=TEAL,align=1)
        equation=(r'$\mathrm{NPY}_{i,t}=\frac{\sum_{j=0}^{11}D_{i,t-j}^{\mathrm{adj}}}{P_{i,t}^{\mathrm{adj}}}'
                  r'+1-\frac{S_{i,t}^{\mathrm{adj}}}{\frac{1}{24}\sum_{j=0}^{23}S_{i,t-j}^{\mathrm{adj}}}$')
        draw_math(pdf,equation,x+20,top-37,width-40,62,size=25)
        paragraph(pdf,'D = adjusted cash dividends; P = adjusted price; S = adjusted shares. Both windows include month t.',
                  x,top-116,width,12.5,color=MUTED)
        notes_h=body_height(notes,width,13.5,6)
        table_top=top-146
        size,layout=fit_table(view['rows'],width,height-146-notes_h-12,17,view['widths'])
        draw_table(pdf,x,table_top,size,layout,prose=True)
        draw_body(pdf,notes,x,bottom+notes_h,width,13.5,6)
        return
    if kind == 'data_cleaning':
        intro_h=body_height(intro,width,15,6)
        footnote=' '.join(view['notes'])
        note_h=paragraph(None,footnote,0,0,width,12,measure=True)
        main_top=top-intro_h-18
        main_h=height-intro_h-note_h-36
        left_w=width*.49
        right_x=x+left_w+24;right_w=width-left_w-24
        size,layout=fit_table(view['rows'],left_w,main_h,view['font'],view['widths'])
        draw_body(pdf,intro,x,top,width,15,6)
        draw_table(pdf,x,main_top,size,layout)
        box_h=(main_h-20)/3
        for i,(heading,detail) in enumerate(view['cleaning']):
            box_top=main_top-i*(box_h+10)
            pdf.setFillColor(STRIPE if i==1 else white);pdf.setStrokeColor(RULE)
            pdf.setLineWidth(.7)
            pdf.roundRect(right_x,box_top-box_h,right_w,box_h,9,stroke=1,fill=1)
            paragraph(pdf,heading,right_x+15,box_top-11,right_w-30,16,bold=True,color=TEAL)
            paragraph(pdf,detail,right_x+15,box_top-39,right_w-30,14,color=INK)
        paragraph(pdf,footnote,x,bottom+note_h,width,12,color=MUTED)
        return
    if kind == 'wide':
        intro_h = body_height(intro, width, 16, 6)
        notes_h = body_height(notes, width, 16, 6)
        gaps = (12 if intro else 0)+(12 if notes else 0)
        extra = view.get('extra')
        if extra:
            # One font size for both tables: the largest at which they fit together.
            for half in range(round(view['font']*2), 25, -1):
                extra_size, extra_layout = fit_table(extra['rows'], width, height, half/2, extra['widths'])
                try:
                    size, layout = fit_table(view['rows'], width, height-intro_h-notes_h-gaps-36-sum(extra_layout[1]),
                                             extra_size, view['widths'])
                except ValueError:
                    continue
                if size == extra_size:
                    break
            else:
                raise ValueError('Slide tables do not fit without changing content')
        else:
            size, layout = fit_table(view['rows'], width, height-intro_h-notes_h-gaps,
                                     view['font'], view['widths'])
        if intro:
            draw_body(pdf, intro, x, top, width, 16, 6)
        table_top = top-intro_h-(12 if intro else 0)
        table_bottom = draw_table(pdf, x, table_top, size, layout, prose=view['prose'])
        if extra:
            paragraph(pdf, extra['caption'], x, table_bottom-12, width, 15, bold=True)
            draw_table(pdf, x, table_bottom-36, extra_size, extra_layout)
        if notes:
            draw_body(pdf, notes, x, bottom+notes_h, width, 16, 6)
        return

    left_w = width*.49
    right_x, right_w = x+left_w+28, width-left_w-28
    if kind == 'with_bottom_table':
        base_size, base_layout = fit_table(tables[0], width, 190, 14, [.20]+[.80/6]*6)
        base_h = sum(base_layout[1])
        notes_h = body_height(notes, width, 13.5, 6)
        main_h = height-base_h-notes_h-24
        size, layout = fit_table(view['rows'], left_w, main_h,
                                 view['font'], view['widths'])
        draw_table(pdf, x, top, size, layout)
        draw_body(pdf, notes, x, bottom+base_h+12+notes_h, width, 13.5, 6)
        draw_picture(pdf, pictures[0], right_x, top, right_w, main_h)
        draw_table(pdf, x, bottom+base_h, base_size, base_layout)
        return

    if kind == 'picture_with_notes_below':
        intro_h = body_height(intro, width, 15, 6)
        notes_h = body_height(notes, width, 15, 6)
        main_h = height-intro_h-notes_h-24
        size, layout = fit_table(view['rows'], left_w, main_h,
                                 view['font'], view['widths'])
        draw_body(pdf, intro, x, top, width, 15, 6)
        main_top = top-intro_h-12
        draw_table(pdf, x, main_top, size, layout)
        draw_picture(pdf, pictures[0], right_x, main_top, right_w, main_h)
        draw_body(pdf, notes, x, bottom+notes_h, width, 15, 6)
        return

    intro_h = body_height(intro, left_w, 16, 6)
    notes_h = body_height(notes, left_w, 14.5, 6)
    gaps = (12 if intro else 0)+(12 if notes else 0)
    size, layout = fit_table(view['rows'], left_w, height-intro_h-notes_h-gaps,
                             view['font'], view['widths'])
    if intro:
        draw_body(pdf, intro, x, top, left_w, 16, 6)
    draw_table(pdf, x, top-intro_h-(12 if intro else 0), size, layout)
    if notes:
        draw_body(pdf, notes, x, bottom+notes_h, left_w, 14.5, 6)
    if kind == 'with_picture':
        draw_picture(pdf, pictures[0], right_x, top, right_w, height)
    elif kind == 'with_side_table':
        rows = [[c.text for c in row.cells] for row in tables[0].rows]
        rows[0] = ['Period', 'Carhart alpha', 'Excess return']
        paragraph(pdf, 'Break-even cost (bps per dollar traded)', right_x, top, right_w, 16, bold=True)
        size_r, layout_r = fit_table(rows, right_w, height-48, 16, [.38, .30, .32])
        draw_table(pdf, right_x, top-45, size_r, layout_r)
    else:
        raise ValueError(f'Unknown numerical layout: {kind}')


def theme_header(pdf, title, number, width, height):
    pdf.setFillColor(CREAM);pdf.rect(0,0,width,height,stroke=0,fill=1)
    pdf.setFillColor(NAVY);pdf.rect(0,height-104,width,104,stroke=0,fill=1)
    pdf.setFillColor(TEAL);pdf.rect(0,height-104,10,104,stroke=0,fill=1)
    section=('STUDY DESIGN' if number==2 else 'DATA AND METHODOLOGY' if number<=5
             else 'US RESULTS' if number<=10 else 'IMPLEMENTATION CHECKS' if number<=12
             else 'PORTFOLIO EXPERIMENTS' if number<=16 else 'CONCLUSIONS' if number==17
             else 'RECENT US RESULTS' if number<=20 else 'UK EXTENSION')
    paragraph(pdf,section,42,height-19,width-84,11,bold=True,color=HexColor('#8BD0C9'))
    paragraph(pdf,title,42,height-44,width-84,30,bold=True,color=white)


def theme_footer(pdf,number,width,dark=False):
    color=HexColor('#B5CACB') if dark else MUTED
    paragraph(pdf,'Yingni Chang and Yikai Shen',42,27,width-120,9,color=color)
    paragraph(pdf,f'{number:02d}',width-76,28,40,10,bold=True,color=color)


def cover(pdf,title,body,width,height):
    pdf.setFillColor(NAVY);pdf.rect(0,0,width,height,stroke=0,fill=1)
    pdf.setFillColor(HexColor('#163F49'));pdf.rect(width-282,0,282,height,stroke=0,fill=1)
    pdf.setFillColor(TEAL);pdf.rect(54,height-74,66,5,stroke=0,fill=1)
    heading=title.paragraphs[0].text.replace('the Conservative','the\nConservative')
    paragraph(pdf,heading,54,height-99,590,45,bold=True,color=white)
    subtitle='\n'.join(p.text for p in list(title.paragraphs)[1:])
    paragraph(pdf,subtitle,57,305,580,12,color=HexColor('#BADEDA'))
    paras=list(body.paragraphs)
    size,gap,_=fit_body(paras,580,191,16,minimum=14)
    top=251
    for p in paras:
        used=paragraph(pdf,p.text,57,top,580,size,color=HexColor('#D1DEDF'))
        top-=used+gap
    x=width-248
    for label,years,top in [('US REPLICATION','1929-2016',391),('RECENT EXTENSION','2017-2024',267)]:
        paragraph(pdf,label,x,top,214,11,bold=True,color=HexColor('#8BD0C9'))
        paragraph(pdf,years,x,top-27,214,28,bold=True,color=white)
    theme_footer(pdf,1,width,True)


def cost_slide(pdf,view,width,height,root):
    theme_header(pdf,view['title'],18,width,height)
    pdf.setFillColor(NAVY);pdf.rect(0,0,244,height-104,stroke=0,fill=1)
    data=results.load('feedback_costs','cost_summary.csv',root=root)[0]
    data=data.loc[data.period.eq('2017–2024')&data.cost_model.eq('flat')].set_index('reference_bps')
    for label,value,top,color in [('Before costs',data.loc[0,'compounded'],386,white),
                                  ('At 30 bps',data.loc[30,'compounded'],257,HexColor('#8BD0C9'))]:
        paragraph(pdf,label,30,top,185,13,color=HexColor('#B6CECE'))
        paragraph(pdf,f'{value*100:.2f}%',30,top-27,185,36,bold=True,color=color)
    paragraph(pdf,'Annual compound return',30,311,185,12,color=HexColor('#B6CECE'))
    pdf.setFillColor(TEAL);pdf.rect(30,280,175,2,stroke=0,fill=1)
    paragraph(pdf,f"Market: {100*data.loc[0,'market_cagr']:.2f}% per year",30,174,185,13,color=white)
    paragraph(pdf,view['intro'],30,138,185,13,color=HexColor('#C5D8D7'))
    rows=[row[:] for row in view['rows']]
    rows[0]=['Cost (bps)','Annual return','Above market / yr','Carhart alpha / yr','Final $100']
    size,layout=fit_table(rows,width-312,290,14.5,[.12,.19,.24,.25,.20])
    draw_table(pdf,278,394,size,layout)
    paragraph(pdf,view['notes'],278,113,width-312,12.5,color=MUTED)
    paragraph(pdf,'Yingni Chang and Yikai Shen',30,27,210,8,color=HexColor('#C5D8D7'))
    paragraph(pdf,'18',width-88,28,40,10,bold=True,color=MUTED)
    pdf.showPage()


def draw_feedback_slide(pdf, view, width, height, root,number):
    if number==18:
        cost_slide(pdf,view,width,height,root);return
    theme_header(pdf,view['title'],number,width,height)
    margin=42;available_width=width-2*margin;top=height-126
    notes=text_paras([view.get('notes','')]);nh=body_height(notes,available_width,14.5,6)
    bottom=48+nh+15
    if 'figure' in view:
        shape=themed_chart(number,root)
        if shape is None:
            path=results.path(*view['figure'],root=root)
            from PIL import Image
            with Image.open(path) as im:iw,ih=im.size
            shape=SimpleNamespace(width=iw,height=ih,image=SimpleNamespace(blob=path.read_bytes()))
        draw_picture(pdf,shape,margin,top,available_width,top-bottom)
    else:
        intro=text_paras([view['intro']]) if view.get('intro') else []
        ih=body_height(intro,available_width,16,6) if intro else 0
        if intro:draw_body(pdf,intro,margin,top,available_width,16,6)
        size,layout=fit_table(view['rows'],available_width,top-bottom-ih-12,17,view['widths'])
        draw_table(pdf,margin,top-ih-(12 if intro else 0),size,layout,prose=view.get('prose',False))
    draw_body(pdf,notes,margin,48+nh,available_width,14.5,6)
    theme_footer(pdf,number,width)
    pdf.showPage()


def export(source, destination, root=None):
    """Build readable PDF evidence tables from the same saved numerical results.

    The PPTX is read only. Five feedback slides follow the original 17 content slides. Existing figures, dimensions and scientific
    results are retained; numerical prose is reorganized into labeled tables.
    """
    deck = Presentation(source)
    width, height = deck.slide_width / EMU, deck.slide_height / EMU
    pdf = canvas.Canvas(str(destination), pagesize=(width, height))
    pdf.setTitle('The Conservative Formula: CRSP Replication, 1929–2024')
    pdf.setAuthor('Yingni Chang and Yikai Shen')
    root = results.RESULTS if root is None else root
    views = make_views(root)
    margin, gutter = 42, 28
    content_width = width-2*margin
    bottom = 40
    fractions = {4: [.36, .64], 6: [.20]+[.80/6]*6,
                 8: [.40, .25, .35], 9: [.17, .09, .20, .27, .27],
                 15: [.25, .19, .23, .15, .18]}
    for number, slide in enumerate(deck.slides, 1):
        if number == 18:
            for extra_number,view in enumerate(extras(root),18):
                draw_feedback_slide(pdf,view,width,height,root,extra_number)
            break  # The old professor-questions slide is excluded from the final submission.
        text_shapes = [s for s in slide.shapes if s.has_text_frame]
        if len(text_shapes) != 3:
            raise ValueError('Expected title, body and author credit on each slide')
        title, body, footer = [s.text_frame for s in text_shapes]
        pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
        tables = [s.table for s in slide.shapes if s.has_table]
        if len(pictures) > 1 or len(tables) > 1:
            raise ValueError('Unexpected slide evidence layout')
        if number==1:
            cover(pdf,title,body,width,height)
            pdf.showPage();continue
        theme_header(pdf,title.paragraphs[0].text,number,width,height)
        restyled=themed_chart(number,root) if pictures else None
        if restyled is not None:pictures=[restyled]
        body_top=height-126
        available = body_top-bottom
        paras = list(body.paragraphs)
        if number == 2:
            research_questions(pdf,margin,body_top,content_width,available)
        elif number == 4:
            selection_flow(pdf,paras,margin,body_top,content_width,available)
        elif number in views:
            draw_numeric_view(pdf, views[number], pictures, tables, margin,
                              body_top, content_width, available)
        elif not pictures and not tables:
            size, gap, used = fit_body(paras, content_width, available, 24)
            draw_body(pdf, paras, margin, body_top-(available-used)*.18,
                      content_width, size, gap)
        elif pictures and tables and number == 9:
            # Keep the rolling chart tall beside the period table and its explanation.
            left_w = content_width*.56
            size_t, layout = fit_table(tables[0], left_w, available*.60, 14,
                                       [.19, .14, .17, .25, .25])
            table_h = sum(layout[1])
            size, gap, used = fit_body(paras, left_w, available-table_h-22, 21)
            draw_body(pdf, paras, margin, body_top, left_w, size, gap)
            draw_picture(pdf, pictures[0], margin+left_w+gutter, body_top,
                         content_width-left_w-gutter, available)
            draw_table(pdf, margin, bottom+table_h, size_t, layout)
        elif pictures and tables:
            # Preserve a wide table below the text/chart comparison.
            table = tables[0]
            size_t, layout = fit_table(table, content_width, 190, 14,
                                       fractions[number])
            table_h = sum(layout[1])
            main_h = available-table_h-20
            left_w = content_width*.49
            size, gap, used = fit_body(paras, left_w, main_h, 21)
            draw_body(pdf, paras, margin, body_top, left_w, size, gap)
            draw_picture(pdf, pictures[0], margin+left_w+gutter, body_top,
                         content_width-left_w-gutter, main_h)
            draw_table(pdf, margin, bottom+table_h, size_t, layout)
        else:
            # Use a balanced text/evidence spread instead of small objects at the top.
            left_w = content_width*(.49 if tables else .46)
            right_w = content_width-left_w-gutter
            size, gap, used = fit_body(paras, left_w, available, 21)
            draw_body(pdf, paras, margin, body_top, left_w, size, gap)
            if pictures:
                draw_picture(pdf, pictures[0], margin+left_w+gutter, body_top,
                             right_w, available)
            else:
                size_t, layout = fit_table(tables[0], right_w, available, 17,
                                           fractions[number])
                table_h = sum(layout[1])
                draw_table(pdf, margin+left_w+gutter, body_top-(available-table_h)*.15,
                           size_t, layout, prose=number == 4)
        theme_footer(pdf,number,width)
        pdf.showPage()
    pdf.save()
    return destination


if __name__ == '__main__':
    source = Path(__file__).resolve().parent / 'conservative_formula.pptx'  # read-only layout input
    output = export(source, Path(__file__).resolve().parent / 'conservative_formula.pdf')
    print(f'Wrote {output}')
