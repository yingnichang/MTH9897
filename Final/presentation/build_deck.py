"""Build the slide deck and speaker notes from the tracked results (results/, through the hash-checked loader).

Project authors: Yingni Chang and Yikai Shen.

    python presentation/build_deck.py            # -> presentation/conservative_formula.pptx, presentation/speaker_notes.md

Every number on a slide is formatted from `research.summary.numbers()` or read from a results table; figures are the
PNGs in results/. Run `python presentation/build_pdf.py` for the requested PDF deliverable.
"""
from __future__ import annotations

import sys
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

FINAL = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(FINAL))
from research import results  # noqa: E402
from research import summary as S  # noqa: E402
from research.summary import ci, num, pct  # noqa: E402

INK, ACCENT, MUTED = RGBColor(0x1F, 0x2A, 0x33), RGBColor(0x17, 0x6B, 0x87), RGBColor(0x6B, 0x75, 0x7D)


class Deck:
    def __init__(self, root: Path):
        self.root = root
        self.prs = Presentation()
        self.prs.core_properties.author = "Yingni Chang and Yikai Shen"
        self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
        self.notes: list[tuple[str, str]] = []

    def slide(self, title: str, bullets=(), notes: str = "", image=None, table: pd.DataFrame | None = None,
              subtitle: str | None = None, text_width: float = 6.2, image_width: float | None = None):
        """Title, then bullets on the left and a figure or table on the right. With both a figure and a table, the
        table spans the full width below the figure. Long bullet lists use a smaller font."""
        s = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        tb = s.shapes.add_textbox(Inches(0.5), Inches(0.3), Inches(12.3), Inches(0.9)).text_frame
        tb.word_wrap = True
        p = tb.paragraphs[0]
        p.text, p.font.size, p.font.bold, p.font.color.rgb = title, Pt(28), True, INK
        if subtitle:
            q = tb.add_paragraph()
            q.text, q.font.size, q.font.color.rgb = subtitle, Pt(15), MUTED
        right = image is not None or table is not None
        width = text_width if right else 12.3
        if bullets:
            body = s.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(width), Inches(5.7)).text_frame
            body.word_wrap = True
            big = 15 if sum(len(b) for b in bullets) / width < 75 else 13
            for i, b in enumerate(bullets):
                para = body.paragraphs[0] if i == 0 else body.add_paragraph()
                level = 1 if b.startswith("  ") else 0
                para.text, para.level = b.strip(), level
                para.font.size, para.font.color.rgb = Pt(big if level == 0 else big - 2), INK if level == 0 else MUTED
                para.space_after = Pt(6)
        left = Inches(0.5 + width + 0.2) if bullets else Inches(0.5)
        avail = Inches(12.3 - (width + 0.2 if bullets else 0))
        top = Inches(1.4)
        if image is not None:
            p = image if isinstance(image, Path) else results.path(*image, root=self.root)
            w = Inches(image_width) if image_width else avail
            pic = s.shapes.add_picture(str(p), left, Inches(1.4), width=w)
            if table is not None:                    # the table goes below the figure, across the slide
                top, left, avail = pic.top + pic.height + Inches(0.15), Inches(0.5), Inches(12.3)
        if table is not None:
            rows, cols = table.shape[0] + 1, table.shape[1] + 1
            t = s.shapes.add_table(rows, cols, left, top, avail, Inches(0.3 * rows)).table
            cells = [[table.index.name or ""] + [str(c) for c in table.columns]] + \
                    [[str(i)] + [str(v) for v in r] for i, r in zip(table.index, table.to_numpy())]
            for r, vals in enumerate(cells):
                for c, v in enumerate(vals):
                    cell = t.cell(r, c)
                    cell.text = v
                    for para in cell.text_frame.paragraphs:
                        para.font.size = Pt(11)
                        para.font.bold = r == 0
        footer = s.shapes.add_textbox(Inches(0.5), Inches(7.15), Inches(12.3), Inches(0.24)).text_frame
        footer.margin_top = footer.margin_bottom = 0
        footer.word_wrap = False
        credit = footer.paragraphs[0]
        credit.text = "Yingni Chang and Yikai Shen"
        credit.font.size, credit.font.color.rgb = Pt(9), MUTED
        s.notes_slide.notes_text_frame.text = notes
        self.notes.append((title, notes))
        return s

    def table_headers(self) -> list[str]:
        """Header cells (first row and first column) of every table on every slide."""
        out = []
        for sl in self.prs.slides:
            for sh in sl.shapes:
                if sh.has_table:
                    rows = list(sh.table.rows)
                    out += [c.text for c in rows[0].cells] + [r.cells[0].text for r in rows]
        return out


def build(root: Path = results.RESULTS, out_dir: Path = FINAL / "presentation") -> Path:
    N = S.numbers(root)
    L = lambda s, n, **k: results.load(s, n, root=root, **k)[0]
    d = Deck(root)
    h, g = "1990–2024|net30", "1990–2024|gross"

    d.slide('Replicating the Conservative Formula',
            subtitle="Yingni Chang and Yikai Shen · MTH 9897 Project 8 · CRSP replication, 1929–2024",
            bullets=["Paper: Blitz and van Vliet (2018), The Conservative Formula, Journal of Portfolio Management 44(7)",
                     "Replicate the US results on the paper's sample, 1929–2016",
                     "Then examine factor exposure, later periods and trading costs",
                     "Main recent comparison: 1990–2024; the paper ends in 2016"],
            notes="We use monthly CRSP data to implement the paper's rule, then extend the backtest through 2024. The first comparison uses the paper's own sample. The later analysis asks whether the return survives factor adjustment and trading costs.")

    d.slide('What we test',
            bullets=["The rule: 1,000 largest US stocks → low-volatility half → rank on 12-1 momentum and net payout yield → best 100, equal weight, quarterly",
                     "1. Known factors: is the return explained by Carhart (primary) or FF5 + momentum (robustness)?",
                     "2. Decay: fixed periods 1929–1962, 1963–1989, 1990–2016, 2017–2024, plus rolling 120-month windows",
                     "3. Costs: net of 10 and 30 bps per dollar traded; break-even cost of the excess return and the alpha",
                     "All estimates with Newey–West (12 lags) or block-bootstrap intervals"],
            notes='We compare the published return and risk first. Then we examine known factors, changes across periods and trading costs. These comparisons were fixed before computing the post-2016 portfolio returns. The additional variants are separate experiments.')

    ex = L("data_audit", "extract_summary.json")
    s = ex["sample"]
    d.slide('The data',
            bullets=[f"CRSP monthly stock file via WRDS, extracted {ex['extracted_at_utc'][:10]}; months {ex['start'][:7]} to {ex['end'][:7]}",
                     f"{s['panel_rows']:,} security-months, {s['securities']:,} securities, {s['terminal_rows']:,} delistings "
                     f"({s['delist_sources']['imputed_shumway']} imputed, {s['delist_sources']['missing_set_zero']} set to zero)",
                     f"Universe: {s['first_formation_universe_n']} stocks in December 1928; the full 1,000 from {s['first_formation_with_full_universe'][:7]}",
                     "One share class per company; dividends implied from total minus price return",
                     "Held stocks without a monthly return: stale-price carry or in-gap price ratio, every case counted",
                     "No release identifier on WRDS: the extract is identified by time stamp and file hashes"],
            notes='The data work matters here. We retain delisted stocks, use historical company and exchange information, and separate missing signal inputs from returns on already-held positions. The early universe is much smaller than 1,000 names.')

    dev = L("replication", "deviations.csv").set_index("item")["ours"]
    d.slide('The portfolio rule',
            bullets=["Paper order: history screen → top 1,000 by market cap → low-volatility half → ranks among complete signals",
                     "Weights drift between rebalances; delisted holdings become cash; costs charged per dollar traded",
                     "Controls: Speculative (high-volatility half, weakest scores), standalone signals, Small, CRSP VW market",
                     "Excluded: Value (no book equity in CRSP) and Table 2 panels C–D (factors not in our sources)"],
            table=dev.to_frame("our convention").head(6).rename_axis("deviation"), text_width=5.6,
            notes='We require 36 months of return history before the size cut. After that, we use the low-volatility half and combine momentum and payout ranks. The largest share class represents each company. We document the choices that the paper leaves open.')

    t1 = L("paper_targets", "table1_vs_paper.csv")
    d.slide('Building the signals',
            bullets=["Volatility: annualized standard deviation of 36 consecutive monthly total returns",
                     "Momentum: compound 11 monthly price returns, excluding the most recent month",
                     "Dividend yield: trailing 12 months of split-adjusted dividends divided by the current adjusted price",
                     "Net payout yield: dividend yield + 1 − latest adjusted shares / average adjusted shares over 24 months",
                     "A falling share count raises payout yield; new share issuance lowers it",
                     "Quarter-end information determines next month's holdings. No forward return enters selection."],
            notes='Momentum uses price returns and excludes the latest month. Portfolio returns include distributions. For payout, both prices and shares need split adjustment: a split should not appear as a repurchase. The 12-month dividend window is our disclosed convention.')

    t1 = t1.loc[t1["row"].isin(["Return (compounded)", "Volatility"]) & t1["status"].eq("compared")]
    t1["paper / ours"] = t1["paper"].map(pct) + " / " + t1["ours"].map(pct)
    tab = t1.pivot(index="row", columns="column", values="paper / ours")[["Formula", "Market", "Small", "Momentum", "Lowvol", "NPY"]]
    d.slide('Our results versus the paper',
            bullets=[f"Conservative {pct(N['rep_cons_compounded'], 2)} compounded, {pct(N['rep_cons_volatility'], 2)} volatility "
                     f"(paper {pct(N['paper_cons_compounded'])}, {pct(N['paper_cons_volatility'])})",
                     f"Market {pct(N['rep_mkt_compounded'], 2)} (paper {pct(N['paper_mkt_compounded'])})",
                     f"Speculative {pct(N['rep_spec_compounded'], 2)} with {pct(N['rep_spec_volatility'])} volatility",
                     f"Conservative beats the market in {N['decades_cons_beats_mkt']} of {N['decades_n']} decades",
                     f"Turnover {pct(N['rep_turnover'])} a quarter vs the paper's {pct(N['paper_turnover'], 0)} (unexplained; measure not stated)"],
            image=("replication", "fig3.png"), table=tab.rename_axis("paper / ours"), text_width=5.0, image_width=6.5,
            notes='The main return and volatility figures are close. The turnover difference is larger and remains unresolved. We have not changed the rule to force agreement. Share classes, the size measure and payout estimation may affect the result, but we cannot separate their effects.')

    d.slide('Factor-adjusted returns',
            bullets=[f"CMS (Conservative − Speculative), 1963–2016: CAPM {pct(N['t2_b_capm_ours'], 2)} vs {pct(N['t2_b_capm_paper'])}; "
                     f"FF5 {pct(N['t2_b_ff5_ours'], 2)} vs {pct(N['t2_b_ff5_paper'])}; FF5 + MOM {pct(N['t2_b_ff5mom_ours'], 2)} vs {pct(N['t2_b_ff5mom_paper'])}",
                     f"1929–2016: CAPM {pct(N['t2_a_capm_ours'], 2)} vs {pct(N['t2_a_capm_paper'])}; Carhart {pct(N['t2_a_carhart_ours'], 2)} vs {pct(N['t2_a_carhart_paper'])} "
                     "(full-sample gap remains)",
                     f"Conservative long-only: Carhart alpha {pct(N['att_cons_carhart_alpha'], 2)} (t {num(N['att_cons_carhart_t'])})",
                     f"  FF5 + momentum, 1963–2016: {pct(N['att_cons_ff5umd_alpha'], 2)} (t {num(N['att_cons_ff5umd_t'])}) — adding profitability and investment reduces alpha",
                     f"CMS net of 30 bps, both legs paying: Carhart {pct(N['att_cms_carhart_net30_alpha'], 2)} (t {num(N['att_cms_carhart_net30_t'])})"],
            notes='The factor model changes the interpretation of the historical return. Adding profitability and investment reduces estimated alpha substantially. The post-1963 comparisons are closer than the full-sample ones; a different sample should not be mistaken for evidence of the cause of the gap.')

    be = L("critical_tests", "breakeven.csv")
    be["bps"] = [f"{b:.0f}" if st == "found" else "none" for b, st in zip(be.breakeven_bps, be.status)]
    bet = be.pivot(index="period", columns="measure", values="bps").rename(columns={
        "net Carhart alpha": "alpha", "net arithmetic excess over the market": "excess return"}).rename_axis("break-even, bps")
    d.slide('How costs change the result',
            bullets=[f"Annual cost estimate: turnover × 8 × cost; quarterly turnover {pct(N['rep_turnover'])} vs {pct(N['paper_turnover'], 0)}",
                     f"Net of 30 bps, 1929–2016: {pct(N['rep_cons_net30_compounded'], 2)} compounded",
                     f"1990–2024 break-even: {N['be_excess|1990–2024']:.0f} bps (excess return), {N['be_alpha|1990–2024']:.0f} bps (alpha)",
                     ("2017–2024: gross excess return and alpha already ≤ 0, so no break-even"
                      if N["be_alpha_status|2017–2024"] == "<= 0" and N["be_excess_status|2017–2024"] == "<= 0"
                      else f"2017–2024 break-even: {N['be_alpha|2017–2024']:.0f} bps (alpha)"),
                     "Early-period break-evens are high, but so were real trading costs then"],
            table=bet, text_width=6.6,
            notes='We charge both purchases and sales. Break-even is the cost per dollar traded at which market excess return or Carhart alpha reaches zero. Flat cost scenarios do not capture historical changes in liquidity or market impact.')

    rows = []
    for per in S.PERIODS:
        rows.append({"period": per, "months": N[f"ct_months|{per}"],
                     "excess, gross": pct(N[f"ct_excess|{per}|gross"]),
                     "Carhart α gross (t)": f"{pct(N[f'ct_alpha|{per}|gross'])} ({num(N[f'ct_alpha_t|{per}|gross'])})",
                     "Carhart α net 30 (t)": f"{pct(N[f'ct_alpha|{per}|net30'])} ({num(N[f'ct_alpha_t|{per}|net30'])})"})
    d.slide('Results by period',
            bullets=[f"Carhart alpha (gross): {pct(N['ct_alpha|1929–1962|gross'])} → {pct(N['ct_alpha|1963–1989|gross'])} → "
                     f"{pct(N['ct_alpha|1990–2016|gross'])} → {pct(N['ct_alpha|2017–2024|gross'])}",
                     "Carhart alpha is statistically significant in the two earlier periods",
                     f"2017–2024 is only {N['ct_months|2017–2024']} months: interval {ci(N['ct_alpha_lo|2017–2024|gross'], N['ct_alpha_hi|2017–2024|gross'])}"],
            image=("critical_tests", "fig_rolling.png"), table=pd.DataFrame(rows).set_index("period"), text_width=7.0,
            image_width=5.0,
            notes='The earlier estimates are stronger than the recent ones. The sample after 2016 is only eight years, so its interval is wide. We do not interpret an insignificant estimate as proof that expected alpha is exactly zero.')

    d.slide('The recent sample: 1990–2024',
            bullets=[f"Gross: excess over market {pct(N['ct_excess|' + g], 2)} {ci(N['ct_excess_lo|' + g], N['ct_excess_hi|' + g])}; "
                     f"Carhart alpha {pct(N['ct_alpha|' + g], 2)} (t {num(N['ct_alpha_t|' + g])})",
                     f"Net of 30 bps: excess {pct(N['ct_excess|' + h], 2)} {ci(N['ct_excess_lo|' + h], N['ct_excess_hi|' + h])}; "
                     f"alpha {pct(N['ct_alpha|' + h], 2)} {ci(N['ct_alpha_lo|' + h], N['ct_alpha_hi|' + h])}",
                     f"FF5 + momentum alpha, net of 30 bps: {pct(N['ct_ff5_alpha|' + h], 2)} (t {num(N['ct_ff5_t|' + h])})",
                     f"Still lower risk: {pct(N['ct_cons_compounded|1990–2024'])} compounded with {pct(N['ct_cons_vol|1990–2024'])} volatility "
                     f"(market {pct(N['ct_mkt_compounded|1990–2024'])}, {pct(N['ct_mkt_vol|1990–2024'])})",
                     f"Cost-inclusive CMS alpha {pct(N['ct_cms_net30_alpha|1990–2024'], 2)} (t {num(N['ct_cms_net30_t|1990–2024'])})"],
            notes="After costs, the recent alpha interval includes zero. The portfolio still has lower volatility than the market. The Conservative-minus-Speculative contrast remains different from the long-only portfolio's performance against the market.")

    pc = L("professor_checks", "payout_comparison.csv")
    payout = pc.loc[pc.period.eq("1929–2016")].set_index("strategy")
    base, div = payout.loc["Baseline NPY"], payout.loc["Dividends-only counterfactual"]
    d.slide('Dividends and share repurchases',
            bullets=["NPY = trailing dividend yield + 1 − latest adjusted shares / 24-month average adjusted shares",
                     "Repurchases raise NPY; share issuance lowers it. Splits must leave the adjusted share count unchanged",
                     f"1929–2016 gross CAGR: {pct(base.compounded, 2)} using NPY versus {pct(div.compounded, 2)} using dividends alone",
                     f"Paired net-30 arithmetic gain from NPY: {pct(base.paired_npy_minus_div_net30_mean, 2)} per year "
                     f"{ci(base.paired_ci_low, base.paired_ci_high, 2)}",
                     "Same universe, low-volatility pool, momentum signal, weights, quarterly dates, and costs in the comparison",
                     "Share-count change measures net repurchases after issuance; it is not reported buyback cash flow"],
            notes='This is the payout issue raised in the assignment. We compare NPY with dividends alone using otherwise identical portfolio rules. The share-count term changes both selections and returns. It measures net repurchases after issuance, rather than reported buyback cash spending.')

    attr = L("professor_checks", "sector_attribution.csv").set_index("period")
    a = attr.loc["1929–2016"]
    quota = payout.loc["Sector quota variant"]
    d.slide('Industry allocation and stock selection',
            bullets=["Paper scores are compared across sectors (page 4); the professor's sector-aware variant is reported separately",
                     f"1929–2016 gross arithmetic excess return: {pct(a.baseline_minus_market_annual, 2)} per year",
                     f"  Industry allocation {pct(a.allocation_annual, 2)} + stock selection {pct(a.selection_annual, 2)} "
                     f"+ reference-to-CRSP-market difference {pct(a.universe_minus_market_annual, 2)}",
                     "Reference: cap-weighted formation universe, quarterly constituents and industries fixed, weights drift, delistings to cash",
                     f"Sector variant: 100 slots proportional to industry security counts; apply the rule separately within each FF12 industry. "
                     f"Gross CAGR {pct(quota.compounded, 2)}, net-30 {pct(quota.net30_compounded, 2)}",
                     "The sector variant is exploratory and changes the paper's global low-volatility screen. Contributions are arithmetic, not causal"],
            notes="We compare with the same formation universe so that industry allocation is not confused with a broad-market universe difference. The attribution includes the selection interaction. We then apply the course's count-proportional quotas, which change the paper's global screen.")

    d.slide('Signal tests',
            bullets=[f"Composite rank IC (1 month): {num(N['ic_comp_h1'], 3)} (NW t {num(N['ic_comp_h1_t'], 1)})",
                     f"Decile spread D10−D1: {pct(N['dec_comp_spread|1963–1989'], 2)} (1963–1989), {pct(N['dec_comp_spread|1990–2016'], 2)} (1990–2016), "
                     f"{pct(N['dec_comp_spread|2017–2024'], 2)} (2017–2024)",
                     f"Factor-mimicking paper portfolio: IR {num(N['fmp_ir_K1'])} at K = 1, {num(N['fmp_ir_K12'])} at K = 12",
                     f"  2017–2024 IR at K = 1: {num(N['fmp_ir_2017_K1'])}",
                     "Low volatility: positive rank IC, flat mean spread (right-skewed high-volatility deciles)"],
            image=("diagnostics", "fig_decay.png"), text_width=5.6,
            notes='These tests use industry-neutral exposures and industry-adjusted returns. Rank IC and a mean decile spread can tell different stories when returns are skewed. The factor-mimicking portfolios are gross diagnostic portfolios and may have substantial leverage.')

    # A single-panel chart keeps the slide's labels legible. Its values come from the verified grid.
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    grid = L("optimization", "grid.csv")
    fig, ax = plt.subplots(figsize=(7, 4.8), layout="constrained")
    for delta, color in [("2%", "#176B87"), ("5%", "#C0504D"), ("none", "#7F7F7F")]:
        for c0, marker in [(10, "o"), (30, "s")]:
            curve = grid.loc[grid.role.eq("spec") & grid.delta.eq(delta) & grid.c0_bps.eq(c0)].sort_values("turnover_one_way_per_quarter")
            ax.plot(100*curve.turnover_one_way_per_quarter, curve.ir, marker=marker, color=color, label=f"Band {delta}; {c0}bps")
    ax.set(xlabel="One-way quarterly turnover (%)", ylabel="Net information ratio", title="Linear optimizer: 1929–2024")
    ax.tick_params(labelsize=11)
    ax.legend(fontsize=10, ncol=2)
    chart = out_dir / "build" / "linear_spec.png"
    chart.parent.mkdir(parents=True, exist_ok=True)
    fig.savefig(chart, dpi=180)
    plt.close(fig)

    d.slide('Long/short and linear optimization',
            bullets=[f"Decile long/short, n = 3, 10 bps, borrow: {pct(N['ls_er'], 2)} a year, IR {num(N['ls_ir'])}, max drawdown {pct(N['ls_mdd'])}",
                     "  Monthly re-weighted cohorts: chosen after the pre-specified buy-and-hold cohorts were ruined "
                     f"({', '.join(x[:7] for x in N['ls_ruin_dates'])})",
                     f"  Capacity 2020s ≈ ${N['ls_capacity_2020s'] / 1e9:.1f}B (scaled to 2024 market size; 10% of volume; no impact model)",
                     f"Linear optimizer, 1990–2024 net 30: Carhart alpha {pct(N['opt_alpha_9024_min'], 2)}–{pct(N['opt_alpha_9024_max'], 2)} "
                     f"(2–5% bands, 20–40% turnover budget); {N['opt_ci_9024_text']}",
                     f"  Unadjusted intervals, selected from a {N['opt_runs_spec']}-run grid; 2017–2024 alphas "
                     f"{pct(N['opt_alpha_1724_min'], 2)} to {pct(N['opt_alpha_1724_max'], 2)} ({N['opt_ci_1724_text']}); "
                     "return-scaled alpha is a later sensitivity",
                     "We also test a portfolio covariance penalty"],
            image=chart, text_width=6.0,
            notes='The initial buy-and-hold long/short loses all capital in the early data. Monthly reweighting was considered after seeing that outcome, and still gives severe drawdowns. The linear optimizer results also need caution: the displayed intervals are unadjusted for testing several settings.')

    qm = L("mean_variance", "summary.csv")
    qr = qm.loc[qm.period.eq("1990–2024")].set_index("risk_aversion")
    qcheck = L("mean_variance", "solver_checks.csv")
    qt = pd.DataFrame({"Net CAGR": qr.compounded.map(lambda x: pct(x, 2)),
                       "Volatility": qr.volatility.map(lambda x: pct(x, 2)),
                       "IR": qr.ir.map(num), "Market corr.": qr.corr_market.map(num)})
    qt.index = [f"γ = {g:g}" for g in qt.index]
    qt.index.name = "1990–2024"
    d.slide('Mean-variance portfolios',
            bullets=["Maximize expected return − half risk aversion × portfolio variance − trading costs",
                     "Covariance: trailing 36 monthly returns, 50% diagonal shrinkage; the window ends at formation",
                     "Return alpha: completed trailing IC × annual volatility × √(3/12) × composite exposure",
                     "Quarterly: 2% stock cap, industry bands ±5pp, two-way turnover budget 40%; scaled costs start at 30bps",
                     f"Two fixed risk aversions (5 and 20); max constraint violation {qcheck.max_constraint_violation.max():.1e}; "
                     f"turnover relaxations {int(qcheck.turnover_relaxations.sum())}",
                     "We chose these settings after examining the baseline; an unseen test is still needed"],
            table=qt, text_width=6.6,
            notes='With only 36 observations for about 1,000 stocks, sample covariance needs regularization. We use half sample covariance and half diagonal covariance. More risk aversion lowers realized volatility, but these long-only portfolios remain highly correlated with the market. We chose these settings after examining the baseline.')

    d.slide('A buffer to reduce trading',
            bullets=["Keep existing holdings while they remain in the eligible pool and within the best 150 ranks",
                     f"Turnover {pct(N['buf_turn_base'])} → {pct(N['buf_turn_buf'])} a quarter ({pct(N['buf_turn_red'])} lower)",
                     f"Net-of-30 paired difference {pct(N['buf_diff'], 2)} a year {ci(N['buf_lo'], N['buf_hi'], 2)}: not significant",
                     f"Gross paired difference {pct(N['buf_gross_diff'], 2)}; the buffer changes holdings as well as costs",
                     "The recent period was already seen when we tested this variant"],
            image=("rank_buffer", "fig_paired_cum.png"), text_width=5.6,
            notes='We fixed the 150-rank buffer before testing it. It reduces turnover, and its net point estimate is positive, but the paired interval includes zero. Since it also changes holdings, we cannot attribute the entire difference to cost savings.')

    d.slide('What we conclude',
            bullets=["Return and volatility are close to the paper; turnover and some regression results differ",
                     "Much of the historical return is associated with known equity factors",
                     f"Alphas decay; 1990–2024 net of 30 bps: alpha {pct(N['ct_alpha|' + h], 2)} {ci(N['ct_alpha_lo|' + h], N['ct_alpha_hi|' + h])}",
                     "The recent net-alpha interval includes zero",
                     "Lower volatility is the most consistent feature of the portfolio",
                     "Limitations: no CRSP release identifier; implied dividends; share-class and size conventions; "
                     "turnover above the paper's; selection across grids; a short post-sample"],
            notes='We reproduce much of the historical return-and-risk pattern. The remaining turnover and regression gaps prevent an exact replication claim. For recent data, lower volatility is more consistent than positive net alpha. The extra strategies need an unseen test before stronger claims would be justified.')

    from professor_questions import TITLE, questions
    d.slide(TITLE, bullets=[f'{topic}: {question}' for topic, question in questions(root)],
            notes='Ask for guidance on payout, industry positioning, turnover, cleaning and the recent extension.')

    bad = results.security_headers(d.table_headers())
    if bad:
        raise ValueError(f"security-level identifier column on a slide: {bad}")
    out_dir.mkdir(parents=True, exist_ok=True)
    from speech_notes import make_notes, plain_notes, write_markdown
    speech = make_notes(root)
    if [n["title"] for n in speech] != [title for title, _ in d.notes]:
        raise ValueError("Speech-note titles must match the slide order")
    for slide, note in zip(d.prs.slides, speech):
        slide.notes_slide.notes_text_frame.text = plain_notes(note)
    path = out_dir / "conservative_formula.pptx"
    d.prs.save(path)
    write_markdown(speech, out_dir / "speaker_notes.md", root)
    return path


if __name__ == "__main__":
    p = build()
    print(f"Wrote {p} and {p.parent / 'speaker_notes.md'}")
